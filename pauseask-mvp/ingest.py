"""Run once: python ingest.py  ->  builds the ChromaDB index from the files in data/.
Expected files (all optional): data/lecture.srt, data/lecture.pdf, data/lecture.pptx, data/unlock_map.json"""
import json, os, re
import chromadb

DATA = "data"
LOCKED = 10**9      # content with no unlock time is hidden from every learner (safe default)
CHUNK_SECONDS = 20  # target length of a transcript chunk
UNLOCK_AT = os.getenv("UNLOCK_AT", "end")  # "end" = strict (no peeking ahead); "start" = a little more context


def to_seconds(s):  # "00:01:05,200" -> 65.2
    h, m, rest = s.strip().split(":")
    sec, ms = rest.replace(".", ",").split(",")
    return int(h) * 3600 + int(m) * 60 + int(sec) + int(ms) / 1000


def read_srt(path):
    out = []
    for block in re.split(r"\n\s*\n", open(path, encoding="utf-8-sig").read().strip()):
        lines = block.strip().splitlines()
        if len(lines) >= 3 and "-->" in lines[1]:
            a, z = lines[1].split("-->")
            out.append((to_seconds(a), to_seconds(z), " ".join(lines[2:])))
    return out


def chunk_srt(lines):
    groups, cur = [], []
    for line in lines:
        cur.append(line)
        if line[1] - cur[0][0] >= CHUNK_SECONDS:
            groups.append(cur); cur = []
    if cur: groups.append(cur)
    return [{"text": " ".join(t for _, _, t in g), "source": "video", "start": g[0][0], "end": g[-1][1],
             "unlock_time": g[-1][1] if UNLOCK_AT == "end" else g[0][0]} for g in groups]


def pdf_chunks(path, umap):
    import fitz  # PyMuPDF
    out = []
    for i, page in enumerate(fitz.open(path), 1):
        text = page.get_text().strip()
        if text:
            out.append({"text": text, "source": "pdf", "page": i, "unlock_time": umap.get("pdf", {}).get(str(i), LOCKED)})
    return out


def ppt_chunks(path, umap):
    from pptx import Presentation
    out = []
    for i, slide in enumerate(Presentation(path).slides, 1):
        text = " ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame).strip()
        if text:
            out.append({"text": text, "source": "ppt", "slide": i, "unlock_time": umap.get("ppt", {}).get(str(i), LOCKED)})
    return out


if __name__ == "__main__":
    umap = json.load(open(f"{DATA}/unlock_map.json")) if os.path.exists(f"{DATA}/unlock_map.json") else {}
    chunks = []
    if os.path.exists(f"{DATA}/lecture.srt"): chunks += chunk_srt(read_srt(f"{DATA}/lecture.srt"))
    if os.path.exists(f"{DATA}/lecture.pdf"): chunks += pdf_chunks(f"{DATA}/lecture.pdf", umap)
    if os.path.exists(f"{DATA}/lecture.pptx"): chunks += ppt_chunks(f"{DATA}/lecture.pptx", umap)
    if not chunks: raise SystemExit("No content found in data/")
    db = chromadb.PersistentClient(path="chroma_data")
    try: db.delete_collection("lecture")
    except Exception: pass
    col = db.create_collection("lecture", metadata={"hnsw:space": "cosine"})  # embeddings are made by Chroma's built-in local model
    col.add(ids=[f"c{i}" for i in range(len(chunks))], documents=[c["text"] for c in chunks],
            metadatas=[{k: v for k, v in c.items() if k != "text"} for c in chunks])
    for s in ("video", "pdf", "ppt"):
        print(s, sum(c["source"] == s for c in chunks), "chunks")
